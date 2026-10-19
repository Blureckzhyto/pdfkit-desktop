"""Versión de la Microsoft Store (MSIX): corre DENTRO de la app empaquetada (tools/msix_check.py
la lanza con Invoke-CommandInDesktopPackage). Comprueba que el paquete se comporte bien:
identidad propia, datos en la carpeta del paquete, carpeta de instalación de solo lectura,
programas incluidos (Poppler, Tesseract, WeasyPrint) ejecutándose desde el paquete, fuentes
privadas, DPI, barra de título oscura y arrastrar y soltar. Fuera del paquete falla a propósito.
Usa documentos generados aquí (nunca archivos del usuario)."""
import ctypes
import os
import subprocess
import sys
import tempfile
from ctypes import wintypes
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from core import paths  # noqa: E402

results, log = [], []


def check(name, cond, detail=""):
    results.append(bool(cond))
    log.append(("  PASS  " if cond else "  FAIL  ") + name + (f"  [{detail}]" if detail else ""))


def attempt(name, func):
    try:
        ok, detail = func()
    except Exception as exc:   # noqa: BLE001
        ok, detail = False, f"{type(exc).__name__}: {exc}"
    check(name, ok, detail)


work = Path(tempfile.mkdtemp(prefix="pliegue_msix_"))
family = paths.package_family_name() or ""
check("corre empaquetada (identidad del paquete)", family.startswith("Blureck.Pliegue_"), family or "sin paquete")

# --- datos y carpeta de instalación
data = paths.data_dir()
expected = Path(os.environ["LOCALAPPDATA"]) / "Packages" / family / "LocalState"
check("datos en %LOCALAPPDATA%\\Packages\\<familia>\\LocalState", data == expected, str(data))


def _writable():
    probe = data / "prueba_escritura.tmp"
    probe.write_text("ok", encoding="utf-8")
    ok = probe.read_text(encoding="utf-8") == "ok"
    probe.unlink()
    return ok, ""


attempt("la carpeta de datos se puede escribir", _writable)
check("la importación de configuración ya corrió al iniciar", (data / paths.IMPORT_MARKER).exists())
check("los registros van a la carpeta de datos", (data / "logs").is_dir(),
      ", ".join(p.name for p in (data / "logs").glob("*")) if (data / "logs").is_dir() else "")
install = Path(sys.executable).parent
check("se ejecuta desde WindowsApps", "windowsapps" in str(install).lower(), str(install))


def _read_only():
    try:
        (install / "no_deberia_poder.tmp").write_text("x", encoding="utf-8")
    except OSError as exc:
        return True, type(exc).__name__
    (install / "no_deberia_poder.tmp").unlink()
    return False, "¡se pudo escribir junto al exe!"


attempt("la carpeta de instalación es de solo lectura", _read_only)
check("la carpeta temporal está fuera de la instalación", not str(work).lower().startswith(str(install).lower()),
      str(work))

# --- identidad en la barra de tareas y DPI
aumid = wintypes.LPWSTR()
hr = ctypes.windll.shell32.GetCurrentProcessExplicitAppUserModelID(ctypes.byref(aumid))
check("no fija un AppUserModelID propio (usa el del paquete)", hr != 0, f"hr={hr & 0xFFFFFFFF:#x}")
awareness = ctypes.c_int(-1)
ctypes.windll.shcore.GetProcessDpiAwareness(None, ctypes.byref(awareness))
check("conciencia de DPI por monitor", awareness.value == 2, str(awareness.value))

# --- programas incluidos, ejecutándose desde el paquete
from core.edit_sign import assets_dir  # noqa: E402

assets = assets_dir()
check("assets dentro de la instalación", str(assets).lower().startswith(str(install).lower()), str(assets))
NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)
for name, exe, args in (("Poppler (pdftoppm)", assets / "poppler" / "pdftoppm.exe", ["-v"]),
                        ("Tesseract", assets / "tesseract" / "tesseract.exe", ["--version"]),
                        ("WeasyPrint", assets / "weasyprint" / "weasyprint.exe", ["--version"])):
    def _helper(exe=exe, args=args):
        proc = subprocess.run([str(exe), *args], capture_output=True, text=True, timeout=120,
                              creationflags=NO_WINDOW)
        text = (proc.stdout + proc.stderr).strip().splitlines()
        return proc.returncode == 0 and bool(text), (text[0] if text else f"código {proc.returncode}")
    attempt(f"{name} se ejecuta desde el paquete", _helper)

# --- herramientas reales con documentos generados
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.pdfgen import canvas  # noqa: E402

sample = work / "muestra.pdf"
c = canvas.Canvas(str(sample), pagesize=A4)
c.setFont("Helvetica-Bold", 28)
c.drawString(72, 760, "Factura de prueba 2026")
c.setFont("Helvetica", 12)
rows = [("Producto", "Cantidad", "Precio"), ("Lapiz", "10", "5.00"), ("Cuaderno", "3", "45.50"),
        ("Regla", "2", "12.00"), ("Goma", "7", "3.25")]
for r, row in enumerate(rows):
    for col, cell in enumerate(row):
        c.drawString(72 + col * 150, 700 - r * 22, cell)
c.save()


def _pdf_to_jpg():
    from core.pdf_tools import pdf_to_images

    out = pdf_to_images(sample, work, dpi=100)
    return len(out) == 1 and out[0].stat().st_size > 5000, out[0].name


def _ocr():
    from core.ocr import ocr_pdf
    from core.pdf_tools import pdf_to_images

    (work / "ocr_img").mkdir()
    image = pdf_to_images(sample, work / "ocr_img", dpi=200)[0]
    scanned = work / "escaneado.pdf"
    from PIL import Image
    Image.open(image).convert("RGB").save(scanned, "PDF", resolution=200)
    out = work / "ocr.pdf"
    ocr_pdf(scanned, out, lang="spa", dpi=200)
    from pypdf import PdfReader
    text = PdfReader(str(out)).pages[0].extract_text() or ""
    return "Factura" in text, text.strip().splitlines()[0] if text.strip() else "(sin texto)"


def _html_file():
    from core.convert_extra import html_to_pdf

    page = work / "pagina.html"
    page.write_text("<html><body><h1>Hola desde Pliegue</h1><p>Prueba HTML a PDF.</p></body></html>",
                    encoding="utf-8")
    out, engine = html_to_pdf(str(page), work / "html.pdf")
    from pypdf import PdfReader
    return "Hola desde Pliegue" in (PdfReader(str(out)).pages[0].extract_text() or ""), engine


def _html_url():
    from core.convert_extra import html_to_pdf

    out, engine = html_to_pdf("https://example.com/", work / "url.pdf")
    from pypdf import PdfReader
    text = PdfReader(str(out)).pages[0].extract_text() or ""
    return "documentation examples" in text, engine


def _pdf_to_excel():
    from core.conversions import pdf_to_excel
    from openpyxl import load_workbook

    out, count = pdf_to_excel(sample, work / "tabla.xlsx")
    values = [str(v) for ws in load_workbook(out).worksheets for row in ws.iter_rows(values_only=True)
              for v in row if v is not None]
    return count >= 1 and any("Cuaderno" in v for v in values), f"{count} tabla(s)"


attempt("PDF a JPG (Poppler)", _pdf_to_jpg)
attempt("OCR (Tesseract, español)", _ocr)
attempt("HTML a PDF desde archivo (WeasyPrint)", _html_file)
attempt("HTML a PDF desde URL (internet)", _html_url)
attempt("PDF a Excel (camelot)", _pdf_to_excel)

# --- ventana: fuentes privadas, barra de título oscura, arrastrar y soltar
import tkinter.font  # noqa: E402

from ui import startup  # noqa: E402
from ui.app import _Root  # noqa: E402

root = _Root()
root.withdraw()
families = set(tkinter.font.families(root))
check("fuentes incluidas registradas (privadas)", "Plus Jakarta Sans" in families or
      any(f.startswith("Plus Jakarta Sans") for f in families))
check("arrastrar y soltar (tkdnd) disponible", getattr(root, "dnd_available", False),
      str(getattr(root, "TkdndVersion", "")))
root.deiconify()
root.update()
hwnd = ctypes.windll.user32.GetParent(root.winfo_id()) or root.winfo_id()
startup.set_dark_titlebar(hwnd, True)
value = ctypes.c_int(0)
ctypes.windll.dwmapi.DwmGetWindowAttribute(ctypes.c_void_p(hwnd), 20, ctypes.byref(value), ctypes.sizeof(value))
check("barra de título oscura", value.value == 1)
root.destroy()

sys.stdout.reconfigure(encoding="utf-8", errors="replace")
print("\n".join(log))
print(f"\n{sum(results)}/{len(results)} OK — {work}")
sys.exit(0 if results and all(results) else 1)
