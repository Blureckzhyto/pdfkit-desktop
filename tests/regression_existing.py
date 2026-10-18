"""Regresión de las herramientas existentes (core)."""
import sys
import tempfile
import traceback
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from openpyxl import Workbook
from PIL import Image
from reportlab.pdfgen import canvas

from pypdf import PdfReader

from core import conversions, image_tools, pdf_tools

TMP = Path(tempfile.mkdtemp(prefix="pliegue_reg_"))
ok = []


def check(name, fn):
    try:
        fn()
        ok.append(True)
        print(f"  PASS  {name}")
    except Exception as exc:
        ok.append(False)
        print(f"  FAIL  {name}: {exc}")
        traceback.print_exc()


def make_pdf(path, pages):
    c = canvas.Canvas(str(path))
    for i in range(pages):
        c.drawString(72, 720, f"Pagina {i + 1}")
        c.showPage()
    c.save()


a, b = TMP / "a.pdf", TMP / "b.pdf"
make_pdf(a, 5)
make_pdf(b, 3)
png, jpg = TMP / "x.png", TMP / "y.jpg"
Image.new("RGB", (300, 200), "red").save(png)
Image.new("RGB", (300, 200), "blue").save(jpg)


def expect(cond, msg="condición no cumplida"):
    if not cond:
        raise AssertionError(msg)


check("unir", lambda: expect(pdf_tools.get_page_count(pdf_tools.merge_pdfs([a, b], TMP / "m.pdf")) == 8))
check("dividir por rangos", lambda: expect(len(pdf_tools.split_by_ranges(a, "1-2, 4", TMP)) == 2))
check("dividir cada N", lambda: expect(len(pdf_tools.split_every_n(a, 2, TMP)) == 3))
check("dividir por tamaño", lambda: expect(len(pdf_tools.split_by_size(a, 0.002, TMP)[0]) >= 2))
check("comprimir", lambda: pdf_tools.compress_pdf(a, TMP / "c.pdf", "medium"))
check("PDF a JPG (Poppler incluido)", lambda: expect(len(pdf_tools.pdf_to_images(b, TMP, dpi=72)) == 3))
check("JPG a PDF", lambda: expect(pdf_tools.get_page_count(pdf_tools.images_to_pdf([png, jpg], TMP / "i.pdf")) == 2))
check("convertir imagen", lambda: image_tools.convert_image(png, TMP, "WEBP"))
check("comprimir imagen", lambda: image_tools.compress_image(jpg, TMP, 60))


def excel():
    wb = Workbook()
    wb.active.append(["A", "B"])
    wb.active.append([1, 2])
    wb.save(TMP / "t.xlsx")
    conversions.excel_to_pdf(TMP / "t.xlsx", TMP / "t.pdf")


check("Excel a PDF", excel)
check("PDF a Word", lambda: conversions.pdf_to_word(a, TMP / "a.docx"))


def pdf_excel():
    """PDF a Excel con una tabla de texto real (camelot; en el exe 1.3.6 fallaba el import)."""
    from openpyxl import load_workbook
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    src = TMP / "tabla.pdf"
    rows = [["Producto", "Cantidad", "Precio"]] + [[f"Articulo {i}", str(i), f"{i * 10}.00"] for i in range(1, 8)]
    table = Table(rows)
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))
    SimpleDocTemplate(str(src), pagesize=letter).build([table])
    out, count = conversions.pdf_to_excel(src, TMP / "tabla.xlsx")
    cells = [str(c.value) for row in load_workbook(out).worksheets[0].iter_rows() for c in row]
    expect(count >= 1 and "Producto" in cells and "Articulo 7" in cells, f"{count} tablas, {cells[:6]}")


check("PDF a Excel (camelot)", pdf_excel)


def word():
    """Word a PDF (Microsoft Word por COM, o el respaldo)."""
    import docx

    document = docx.Document()
    document.add_paragraph("Hola desde Word")
    document.save(TMP / "w.docx")
    out = conversions.word_to_pdf(TMP / "w.docx", TMP / "w.pdf")
    expect(Path(out).exists() and "Hola desde Word" in PdfReader(str(out)).pages[0].extract_text())


check("Word a PDF", word)
print(f"\n{sum(ok)}/{len(ok)} OK")
sys.exit(0 if all(ok) else 1)
