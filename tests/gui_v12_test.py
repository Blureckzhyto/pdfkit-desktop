"""Prueba de punta a punta de la versión 1.2 (vistas nuevas + regresión) bajo mainloop."""
import sys
import tempfile
import time
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import img2pdf
from PIL import Image, ImageDraw, ImageFont, ImageGrab
from pypdf import PdfReader
from reportlab.pdfgen import canvas

SHOTS = Path(__file__).parent / "_capturas" / "shots_v12"
SHOTS.mkdir(parents=True, exist_ok=True)
SHOTS.mkdir(exist_ok=True)
TMP = Path(tempfile.mkdtemp(prefix="pliegue_gui12_"))
SETTINGS = __import__("core.paths").paths.settings_file()
backup = SETTINGS.read_bytes() if SETTINGS.exists() else None


def make_pdf(path, pages):
    c = canvas.Canvas(str(path), pagesize=(595, 842))
    for lines in pages:
        c.setFont("Helvetica", 18)
        for n, line in enumerate(lines):
            c.drawString(72, 760 - n * 26, line)
        c.showPage()
    c.save()


doc = TMP / "doc.pdf"
make_pdf(doc, [[f"PAGINA{i + 1}", "Texto de ejemplo para las pruebas"] for i in range(6)])
doc_b = TMP / "doc_b.pdf"
make_pdf(doc_b, [["PAGINA1", "Texto CAMBIADO para las pruebas"]] + [[f"PAGINA{i + 1}", "Texto de ejemplo para las pruebas"] for i in range(1, 6)])
scan = TMP / "escaneo.pdf"
img = Image.new("RGB", (1700, 2200), "white")
ImageDraw.Draw(img).text((150, 200), "Factura de servicios", fill="black",
                         font=ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 70))
buf = BytesIO()
img.save(buf, "PNG", dpi=(200, 200))
scan.write_bytes(img2pdf.convert(buf.getvalue()))
html = TMP / "pagina.html"
html.write_text("<html><meta charset='utf-8'><body><h1>Reporte anual</h1></body></html>", encoding="utf-8")
broken = TMP / "roto.pdf"
raw = doc.read_bytes()
broken.write_bytes(raw[:raw.rfind(b"xref")] + b"xref\nbasura\n%%EOF\n")
from pptx import Presentation
prs = Presentation()
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = "Plan de proyecto"
deck = TMP / "deck.pptx"
prs.save(deck)

from ui import theme
from ui.app import App

app = App()
app.attributes("-topmost", True)
results, log = [], []


def check(name, cond, detail=""):
    results.append(bool(cond))
    log.append(("  PASS  " if cond else "  FAIL  ") + name + (f"  [{detail}]" if detail else ""))


def grab(name):
    app.update_idletasks()
    app.update()
    x, y = app.winfo_rootx(), app.winfo_rooty()
    ImageGrab.grab(bbox=(x, y, x + app.winfo_width(), y + app.winfo_height())).save(SHOTS / f"{name}.png")


class E:
    def __init__(self, x, y):
        self.x, self.y = x, y


def view(key):
    app.show_view(key)
    v = app._views[key]
    if getattr(v, "output_selector", None):
        v.output_selector.entry.delete(0, "end")
        v.output_selector.entry.insert(0, str(TMP))
    return v


def status(v):
    return v.status.cget("text")


# --------------------------------------------------------------- secuencia de pasos
steps = []


def step(fn):
    steps.append(fn)
    return fn


def run_steps():
    if not steps:
        finish()
        return
    fn = steps.pop(0)
    try:
        wait_for = fn()
    except Exception as exc:
        import traceback
        traceback.print_exc()
        check(f"{fn.__name__}: sin excepciones", False, repr(exc))
        wait_for = None
    deadline = time.time() + 120

    def poll():
        if wait_for is None or wait_for() or time.time() > deadline:
            app.after(80, run_steps)
        else:
            app.after(120, poll)

    app.after(150, poll)


def idle(v):
    return lambda: not v._running


@step
def s_sidebar():
    cats = [c for c, _color, _t in app.categories]
    check("categorías nuevas en orden", cats[:6] == ["ORDENAR PDF", "OPTIMIZAR PDF", "CONVERTIR A PDF",
                                                     "CONVERTIR DESDE PDF", "EDITAR PDF", "EDITAR & FIRMAR"], str(cats))
    names = [n for _c, _col, tools in app.categories for _k, _s, n, _v in tools]
    expected_new = ["Extraer páginas", "Ordenar páginas", "Reparar PDF", "OCR PDF", "PowerPoint a PDF",
                    "HTML a PDF", "PDF a PowerPoint", "PDF a PDF/A", "Rotar PDF", "Números de página",
                    "Recortar PDF", "Comparar PDF"]
    check("las 12 herramientas nuevas están en la barra", all(n in names for n in expected_new))
    check("herramientas existentes conservadas", all(n in names for n in [
        "Unir PDF", "Dividir PDF", "Comprimir PDF", "JPG a PDF", "Word a PDF", "Excel a PDF", "PDF a JPG",
        "PDF a Word", "PDF a Excel", "Marca de agua", "Editar PDF", "Firmar PDF", "Convertir imágenes",
        "Comprimir imágenes"]))
    import customtkinter as ctk
    # desde 1.3.1 hay tema oscuro: el modo debe coincidir con la preferencia guardada
    check("tema según la preferencia (claro/oscuro/sistema)",
          ctk.get_appearance_mode() == ("Dark" if app._mode == "dark" else "Light"), app.theme_pref)


@step
def s_extract():
    v = view("extract")
    v.picker.add_files([str(doc)])
    v.range_entry.insert(0, "6, 2-3")
    return lambda: v.page_count == 6


@step
def s_extract_run():
    v = app._views["extract"]
    check("extraer: muestra la cantidad de páginas", "6 página" in v.count_label.cget("text"))
    v.range_entry.delete(0, "end")
    v.range_entry.insert(0, "1, 9")
    v.run()
    check("extraer: valida el rango antes de procesar", "fuera del documento" in status(v) and not v._running)
    v.range_entry.delete(0, "end")
    v.range_entry.insert(0, "6, 2-3")
    v.run()
    return idle(v)


@step
def s_extract_check():
    out = TMP / "doc_extraido.pdf"
    order = [p.extract_text().split("\n")[0] for p in PdfReader(str(out)).pages] if out.exists() else []
    check("extraer: PDF con páginas 6, 2, 3", order == ["PAGINA6", "PAGINA2", "PAGINA3"], str(order))
    v = view("delete")
    v.picker.add_files([str(doc)])
    v.range_entry.insert(0, "1, 3")
    return lambda: v.page_count == 6


@step
def s_delete():
    v = app._views["delete"]
    v.run()
    return idle(v)


@step
def s_organize():
    out = TMP / "doc_sin_paginas.pdf"
    check("eliminar páginas", out.exists() and len(PdfReader(str(out)).pages) == 4, status(app._views["delete"]))
    v = view("organize")
    v.picker.add_files([str(doc)])
    # las miniaturas se generan a medida que se ven (1.3.1): las 6 caben en pantalla
    return lambda: len(v.order) == 6 and len(v._photos) == 6


@step
def s_organize_drag():
    v = app._views["organize"]
    check("ordenar: 6 miniaturas", len(v.order) == 6 and len(v._photos) == 6)
    # arrastrar la página 1 detrás de la 3
    x1, y1 = v._cell_origin(0)
    x3, y3 = v._cell_origin(2)
    v._on_press(E(x1 + 40, y1 + 40))
    v._on_drag(E(x1 + 80, y1 + 60))
    v._on_drag(E(x3 + 120, y3 + 60))
    v._on_release(E(x3 + 120, y3 + 60))
    check("ordenar: arrastrar reordena", v.order[:3] == [1, 2, 0], str(v.order))
    v.selected = v.order.index(5)
    v._move(-1)
    v._delete()
    check("ordenar: mover y eliminar", len(v.order) == 5 and 5 not in v.order, str(v.order))
    grab("gui_organize")
    v.run()
    return idle(v)


@step
def s_repair():
    out = TMP / "doc_ordenado.pdf"
    firsts = [p.extract_text().split("\n")[0] for p in PdfReader(str(out)).pages] if out.exists() else []
    check("ordenar: PDF guardado en el nuevo orden", firsts[:3] == ["PAGINA2", "PAGINA3", "PAGINA1"], str(firsts))
    v = view("repair")
    v.picker.add_files([str(broken)])
    v.run()
    return idle(v)


@step
def s_ocr():
    v = app._views["repair"]
    report = v.results.get("1.0", "end")
    check("reparar: copia guardada e informe", status(v).startswith("✓") and "corregido" in report, status(v))
    v = view("ocr")
    v.picker.add_files([str(scan)])
    v.run()
    return idle(v)


@step
def s_pptx():
    v = app._views["ocr"]
    out = TMP / "escaneo_ocr.pdf"
    check("OCR: texto buscable", out.exists() and "Factura" in PdfReader(str(out)).pages[0].extract_text(), status(v))
    v = view("pptx2pdf")
    v.picker.add_files([str(deck)])
    v.run()
    return idle(v)


@step
def s_html():
    v = app._views["pptx2pdf"]
    out = TMP / "deck.pdf"
    check("PowerPoint a PDF", out.exists() and "Plan de proyecto" in PdfReader(str(out)).pages[0].extract_text(), status(v))
    v = view("html2pdf")
    v.picker.add_files([str(html)])
    v.run()
    return idle(v)


@step
def s_pdf2pptx():
    v = app._views["html2pdf"]
    out = TMP / "pagina.pdf"
    check("HTML a PDF", out.exists() and "Reporte" in PdfReader(str(out)).pages[0].extract_text(), status(v))
    v = view("pdf2pptx")
    v.picker.add_files([str(doc)])
    v.dpi.set("100")
    v.run()
    return idle(v)


@step
def s_pdfa():
    v = app._views["pdf2pptx"]
    check("PDF a PowerPoint", (TMP / "doc.pptx").exists() and "6 diapositiva" in status(v), status(v))
    v = view("pdfa")
    v.picker.add_files([str(doc)])
    v.run()
    return idle(v)


@step
def s_rotate():
    v = app._views["pdfa"]
    check("PDF a PDF/A informa el nivel", "PDF/A-1b" in status(v) and (TMP / "doc_pdfa.pdf").exists(), status(v)[:90])
    v = view("rotate")
    v.picker.add_files([str(doc)])
    return lambda: len(v.menus) == 6


@step
def s_rotate_run():
    v = app._views["rotate"]
    v.run()
    check("rotar: avisa si no hay cambios", "No cambiaste" in status(v))
    v.all_angle.set("180°")
    v._apply_all()
    v.menus[0].set("90°")
    v.run()
    return idle(v)


@step
def s_numbers():
    out = TMP / "doc_rotado.pdf"
    rots = [p.rotation for p in PdfReader(str(out)).pages] if out.exists() else []
    check("rotar: por página y todas", rots == [90, 180, 180, 180, 180, 180], str(rots))
    v = view("numbers")
    v.picker.add_files([str(doc)])
    v.fmt.set("Página 1 de N")
    v.position.set("Pie derecho")
    v.run()
    return idle(v)


@step
def s_crop():
    out = TMP / "doc_numerado.pdf"
    check("números de página", out.exists() and "Página 6 de 6" in PdfReader(str(out)).pages[5].extract_text(),
          status(app._views["numbers"]))
    v = view("crop")
    v.picker.add_files([str(doc)])
    return lambda: v._view is not None


@step
def s_crop_drag():
    v = app._views["crop"]
    x0, y0, x1, y1 = v._crop_px()
    v._on_press(E((x0 + x1) / 2, y0))
    v._on_drag(E((x0 + x1) / 2, y0 + 40))
    v._drag_edge = None
    top_mm = float(v.entries["top"].get())
    check("recortar: arrastrar el borde superior actualiza el margen", top_mm > 10, f"{top_mm} mm")
    v.entries["left"].delete(0, "end")
    v.entries["left"].insert(0, "15")
    v._entry_changed("left")
    check("recortar: escribir el margen mueve el borde", abs(v.margins["left"] - 15 * 72 / 25.4) < 0.01)
    grab("gui_crop")
    v.run()
    return idle(v)


@step
def s_compare_text():
    out = TMP / "doc_recortado.pdf"
    w = float(PdfReader(str(out)).pages[0].mediabox.width) if out.exists() else 0
    check("recortar: PDF recortado", abs(w - (595 - 15 * 72 / 25.4)) < 0.2, f"ancho {w:.1f}")
    v = view("compare")
    v.pickers[0].add_files([str(doc)])
    v.pickers[1].add_files([str(doc_b)])
    v.run()
    return idle(v)


@step
def s_compare_visual():
    v = app._views["compare"]
    left = v.texts[0].get("1.0", "end")
    right = v.texts[1].get("1.0", "end")
    check("comparar texto: línea quitada y agregada", "− Texto de ejemplo" in left and "+ Texto CAMBIADO" in right,
          status(v))
    grab("gui_compare")
    v.mode.set("VISUAL")
    v._mode_changed("VISUAL")
    v.run()
    return idle(v)


@step
def s_censor():
    v = app._views["compare"]
    check("comparar visual: detecta la página 1", "1 página(s) con diferencias" in status(v), status(v))
    v = view("redact")
    check("censurar: solo herramienta REDACTAR", list(v.tool.cget("values")) == ["REDACTAR"])
    v.picker.add_files([str(doc)])
    return lambda: v._base_photo is not None


@step
def s_censor_run():
    v = app._views["redact"]
    s = 100 / 72
    v._on_press(E(60 * s, 60 * s))
    v._on_drag(E(200 * s, 90 * s))
    v._on_release(E(200 * s, 90 * s))
    v.run()
    return idle(v)


@step
def s_protect():
    out = TMP / "doc_censurado.pdf"
    t = PdfReader(str(out)).pages[0].extract_text() if out.exists() else "?"
    check("censurar: borra el texto debajo", "PAGINA1" not in t and "Texto de ejemplo" in t, repr(t[:40]))
    v = view("protect")
    check("proteger: sin selector de modo", not v.mode.winfo_ismapped())
    v.picker.add_files([str(doc)])
    v.user_pw.insert(0, "clave")
    v.run()
    return idle(v)


@step
def s_unprotect():
    out = TMP / "doc_protegido.pdf"
    check("proteger PDF", out.exists() and PdfReader(str(out)).is_encrypted, status(app._views["protect"]))
    v = view("unprotect")
    v.picker.add_files([str(out)])
    v.unlock_pw.insert(0, "clave")
    v.run()
    return idle(v)


@step
def s_regression():
    check("desproteger PDF", (TMP / "doc_protegido_desbloqueado.pdf").exists(), status(app._views["unprotect"]))
    v = view("merge")
    v.picker.add_files([str(doc), str(doc_b)])
    v.run()
    return idle(v)


@step
def s_regression2():
    check("regresión: unir PDF", (TMP / "doc_unido.pdf").exists(), status(app._views["merge"]))
    v = view("split")
    v.picker.add_files([str(doc)])
    v.mode.set("Tamaño máximo (MB)")
    v._mode_changed("Tamaño máximo (MB)")
    v.value_entry.delete(0, "end")
    v.value_entry.insert(0, "0.001")
    v.run()
    return idle(v)


@step
def s_regression3():
    v = app._views["split"]
    check("regresión: dividir por tamaño", "archivo(s)" in status(v), status(v)[:80])
    v = view("watermark")
    v.picker.add_files([str(doc)])
    v.run()
    return idle(v)


@step
def s_done():
    v = app._views["watermark"]
    check("regresión: marca de agua", status(v).startswith("✓"), status(v)[:80])
    check("barra de estado con color de éxito", v.status_bar.cget("fg_color") == theme.SUCCESS_BG)
    grab("gui_watermark")
    app.show_view("home")


def finish():
    grab("gui_home_end")
    app.destroy()


app.after(800, run_steps)
app.after(600000, app.destroy)
try:
    app.mainloop()
finally:
    if backup is not None:
        SETTINGS.write_bytes(backup)
print("\n".join(log))
print(f"\n{sum(results)}/{len(results)} OK — {TMP}")
sys.exit(0 if results and all(results) else 1)
