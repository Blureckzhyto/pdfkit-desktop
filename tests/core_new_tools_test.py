"""Pruebas de las 12 herramientas nuevas (core)."""
import sys
import tempfile
import traceback
from io import BytesIO
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader
from reportlab.pdfgen import canvas

from core import compare, convert_extra, ocr, page_tools

TMP = Path(tempfile.mkdtemp(prefix="pliegue_new_"))
results = []


def check(name, cond, detail=""):
    results.append(bool(cond))
    print(("  PASS  " if cond else "  FAIL  ") + name + (f"  [{detail}]" if detail else ""))


def run(name, fn):
    print(f"\n[{name}]")
    try:
        fn()
    except Exception:
        results.append(False)
        print(f"  FAIL  {name}: excepción")
        traceback.print_exc()


def text(path, i=None):
    r = PdfReader(str(path))
    pages = r.pages if i is None else [r.pages[i]]
    return "\n".join(p.extract_text() or "" for p in pages)


def make_pdf(path, lines_per_page):
    c = canvas.Canvas(str(path), pagesize=(595, 842))
    for lines in lines_per_page:
        c.setFont("Helvetica", 16)
        for n, line in enumerate(lines):
            c.drawString(72, 760 - n * 24, line)
        c.showPage()
    c.save()


doc = TMP / "doc.pdf"
make_pdf(doc, [[f"PAGINA{i + 1}"] for i in range(5)])


def t_pages():
    out = TMP / "extract.pdf"
    _, n = page_tools.extract_pages(doc, out, "3, 1-2")
    order = [l.strip() for l in text(out).split("\n") if l.strip()]
    check("extraer respeta el orden escrito", n == 3 and order == ["PAGINA3", "PAGINA1", "PAGINA2"], str(order))
    try:
        page_tools.extract_pages(doc, TMP / "x.pdf", "1, 9")
        check("extraer valida el rango antes de procesar", False)
    except ValueError as exc:
        check("extraer valida el rango antes de procesar", "fuera del documento" in str(exc), str(exc))
    _, removed, left = page_tools.delete_pages(doc, TMP / "del.pdf", "2, 4")
    check("eliminar páginas", removed == 2 and left == 3 and "PAGINA2" not in text(TMP / "del.pdf"))
    try:
        page_tools.delete_pages(doc, TMP / "x.pdf", "1-5")
        check("no deja eliminar todas", False)
    except ValueError:
        check("no deja eliminar todas", True)
    page_tools.reorder_pages(doc, TMP / "ord.pdf", [4, 3, 0])
    check("ordenar páginas", [l for l in text(TMP / "ord.pdf").split("\n") if l] == ["PAGINA5", "PAGINA4", "PAGINA1"])
    thumbs = page_tools.render_thumbnails(doc, dpi=30)
    check("miniaturas de todas las páginas", len(thumbs) == 5 and thumbs[0].width < 300, f"{len(thumbs)}, {thumbs[0].size}")


def t_rotate_numbers_crop():
    out, changed = page_tools.rotate_pages(doc, TMP / "rot.pdf", {0: 90, 1: 180, 2: 0})
    check("rotar por página", page_tools.get_rotations(out)[:3] == [90, 180, 0] and changed == 2,
          str(page_tools.get_rotations(out)))
    out, n = page_tools.add_page_numbers(doc, TMP / "num.pdf", "bottom_right", start=3,
                                         fmt="pagina_n_de_total", size=10, color="#E53935")
    check("números: formato 'Página X de N' con inicio 3", "Página 3 de 7" in text(out, 0)
          and "Página 7 de 7" in text(out, 4), text(out, 0)[-20:].strip())
    out, n = page_tools.add_page_numbers(doc, TMP / "num2.pdf", "top_center", fmt="n_de_total",
                                         pages_spec="2-3")
    check("números: rango de páginas", n == 2 and "1 / 2" in text(out, 1) and "/" not in text(out, 0))
    out, n = page_tools.crop_pages(doc, TMP / "crop.pdf", (10, 20, 10, 10), pages=[0])
    p = PdfReader(str(out)).pages
    mm = 72 / 25.4
    check("recortar: márgenes en mm", abs(float(p[0].mediabox.width) - (595 - 20 * mm)) < 0.1
          and abs(float(p[0].mediabox.height) - (842 - 30 * mm)) < 0.1
          and p[0].cropbox == p[0].mediabox, f"{float(p[0].mediabox.width):.1f}x{float(p[0].mediabox.height):.1f}")
    check("recortar solo la página actual", float(p[1].mediabox.width) == 595)
    try:
        page_tools.crop_pages(doc, TMP / "x.pdf", (200, 200, 0, 0))
        check("recortar valida márgenes", False)
    except ValueError:
        check("recortar valida márgenes", True)


def t_repair():
    raw = doc.read_bytes()
    broken = TMP / "roto.pdf"
    # destruir la tabla de referencias cruzadas
    cut = raw.rfind(b"xref")
    broken.write_bytes(raw[:cut] + b"xref\ngarbage garbage\n%%EOF\n")
    out, warnings, pages = page_tools.repair_pdf(broken, TMP / "reparado.pdf")
    check("reparar: recupera las páginas", pages == 5 and "PAGINA5" in text(out), f"{pages} páginas")
    check("reparar: informa los problemas encontrados", len(warnings) > 0, warnings[0][:70] if warnings else "")
    junk = TMP / "basura.pdf"
    junk.write_bytes(b"esto no es un pdf" * 50)
    try:
        page_tools.repair_pdf(junk, TMP / "x.pdf")
        check("reparar: error claro si es irrecuperable", False)
    except RuntimeError as exc:
        check("reparar: error claro si es irrecuperable", "No se pudo recuperar" in str(exc), str(exc)[:80])


def t_ocr():
    scan = TMP / "escaneo.pdf"
    img = Image.new("RGB", (1700, 2200), "white")
    font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 64)
    d = ImageDraw.Draw(img)
    d.text((150, 200), "Contrato de alquiler", fill="black", font=font)
    d.text((150, 320), "Pago mensual acordado", fill="black", font=font)
    buf = BytesIO()
    img.save(buf, "PNG", dpi=(200, 200))
    import img2pdf
    scan.write_bytes(img2pdf.convert(buf.getvalue()))
    check("el escaneo no tiene texto", not text(scan).strip())
    progress_msgs = []
    out, done, skipped = ocr.ocr_pdf(scan, TMP / "ocr.pdf", lang="spa",
                                     progress=lambda f, m=None: m and progress_msgs.append(m))
    t = text(out)
    check("OCR agrega texto buscable (español)", "Contrato" in t and "alquiler" in t, repr(t[:60]))
    check("OCR informa progreso por página", "Procesando página 1 de 1…" in progress_msgs)
    check("OCR conserva la imagen original", abs(Path(out).stat().st_size - scan.stat().st_size) < 120_000,
          f"{scan.stat().st_size} -> {Path(out).stat().st_size} bytes")
    texty = TMP / "con_texto.pdf"
    make_pdf(texty, [["Este documento ya tiene texto seleccionable de verdad."]] * 3)
    out2, done2, skipped2 = ocr.ocr_pdf(texty, TMP / "ocr2.pdf", lang="spa+eng")
    check("OCR saltea páginas que ya tienen texto", done2 == 0 and skipped2 == 3, f"{done2}/{skipped2}")


def t_pptx():
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Informe trimestral"
    s.placeholders[1].text = "Ventas en alza"
    pic = TMP / "pic.png"
    Image.new("RGB", (200, 100), (30, 136, 229)).save(pic)
    s2 = prs.slides.add_slide(prs.slide_layouts[6])
    s2.shapes.add_picture(str(pic), Inches(1), Inches(1), Inches(3))
    box = s2.shapes.add_textbox(Inches(1), Inches(4), Inches(6), Inches(1))
    box.text_frame.text = "Segunda diapositiva"
    box.text_frame.paragraphs[0].runs[0].font.size = Pt(28)
    src = TMP / "deck.pptx"
    prs.save(src)
    out, method, note = convert_extra.pptx_to_pdf(src, TMP / "deck.pdf")
    t = text(out)
    check("PowerPoint a PDF", len(PdfReader(str(out)).pages) == 2 and "Informe trimestral" in t,
          f"método: {method}")
    basic = TMP / "deck_basic.pdf"
    convert_extra._pptx_basic(src, basic)
    tb = text(basic)
    check("renderizador básico (python-pptx + ReportLab)", "Informe trimestral" in tb and "Segunda diapositiva" in tb)
    out, n = convert_extra.pdf_to_pptx(doc, TMP / "doc.pptx", dpi=60)
    back = Presentation(str(out))
    check("PDF a PowerPoint: una diapositiva por página", n == 5 and len(back.slides) == 5
          and back.slides[0].shapes[0].shape_type == 13)


def t_html():
    page = TMP / "pagina.html"
    page.write_text("<html><head><meta charset='utf-8'></head><body><h1>Factura Nº 42</h1>"
                    "<p>Total: $ 1.500</p></body></html>", encoding="utf-8")
    out, engine = convert_extra.html_to_pdf(str(page), TMP / "html.pdf")
    check("HTML a PDF con WeasyPrint incluido", engine == "WeasyPrint" and "Factura" in text(out), engine)
    original = convert_extra.assets_dir
    convert_extra.assets_dir = lambda: TMP / "no_existe"
    try:
        out, engine = convert_extra.html_to_pdf(str(page), TMP / "html_edge.pdf")
    finally:
        convert_extra.assets_dir = original
    check("HTML a PDF: respaldo con Edge", engine == "Microsoft Edge" and "Factura" in text(out), engine)
    try:
        convert_extra.html_to_pdf(str(TMP / "nada.html"), TMP / "x.pdf")
        check("HTML: archivo inexistente", False)
    except ValueError:
        check("HTML: archivo inexistente", True)


def t_pdfa():
    import pikepdf
    out, level = convert_extra.pdf_to_pdfa(doc, TMP / "pdfa.pdf")
    with pikepdf.open(out) as pdf:
        meta = pdf.open_metadata()
        intents = pdf.Root.get("/OutputIntents")
        check("PDF/A: metadatos pdfaid 1B", meta.get("pdfaid:part") == "1" and meta.get("pdfaid:conformance") == "B",
              f"{meta.get('pdfaid:part')}{meta.get('pdfaid:conformance')}")
        check("PDF/A: OutputIntent con perfil ICC", intents is not None and "/DestOutputProfile" in intents[0])
        check("PDF/A: versión 1.4 y 5 páginas", pdf.pdf_version == "1.4" and len(pdf.pages) == 5, pdf.pdf_version)
    check("PDF/A: informa el nivel alcanzado", "básico" in level or "completo" in level, level[:40])


def t_compare():
    a, b = TMP / "a.pdf", TMP / "b.pdf"
    make_pdf(a, [["Linea uno", "Linea dos", "Linea tres"], ["Igual"]])
    make_pdf(b, [["Linea uno", "Linea DOS cambiada", "Linea tres", "Linea nueva"], ["Igual"], ["Extra"]])
    res = compare.compare_text(a, b)
    states = [(r[1], r[3]) for r in res["rows"]]
    check("comparar texto: detecta cambios", res["removed"] == 1 and res["added"] == 3 and not res["identical"],
          f"-{res['removed']} +{res['added']}")
    check("comparar texto: filas alineadas", ("removed", "added") in states and ("equal", "equal") in states)
    check("comparar texto: cantidad de páginas distinta", res["pages_a"] == 2 and res["pages_b"] == 3)
    vis = compare.compare_visual(a, b, TMP / "visual.pdf", dpi=50)
    check("comparar visual: página 1 cambia, página 2 igual", vis["changed_pages"] == [1], str(vis["changed_pages"]))
    check("comparar visual: genera el informe PDF", len(PdfReader(str(vis["path"])).pages) == 2)
    same = compare.compare_text(a, a)
    check("comparar: PDFs idénticos", same["identical"])


for name, fn in [("ordenar", t_pages), ("rotar / números / recortar", t_rotate_numbers_crop),
                 ("reparar", t_repair), ("OCR", t_ocr), ("PowerPoint", t_pptx), ("HTML", t_html),
                 ("PDF/A", t_pdfa), ("comparar", t_compare)]:
    run(name, fn)
print(f"\n{sum(results)}/{len(results)} OK — {TMP}")
sys.exit(0 if all(results) else 1)
