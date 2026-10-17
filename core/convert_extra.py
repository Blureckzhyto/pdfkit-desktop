"""Conversiones adicionales: PowerPoint <-> PDF, HTML -> PDF y PDF -> PDF/A.
Sin imports de UI; mensajes en español."""
from __future__ import annotations

import glob
import os
import shutil
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path

from .edit_sign import _noop, _open, assets_dir

EMU_PER_PT = 12700
_NO_WINDOW = getattr(subprocess, "CREATE_NO_WINDOW", 0)
BASIC_PPTX_NOTE = ("Conversión básica: los diseños complejos, degradados, fuentes "
                   "especiales y efectos pueden no reproducirse exactamente.")


def _run(cmd: list[str], timeout: int = 300) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, capture_output=True, timeout=timeout, creationflags=_NO_WINDOW)


def _last_error_line(proc) -> str:
    text = (proc.stderr or b"").decode("utf-8", "replace").strip().splitlines()
    errors = [line for line in text if "error" in line.lower()] or text
    return errors[-1] if errors else f"código {proc.returncode}"


# --------------------------------------------------------------------------- PowerPoint -> PDF

def pptx_to_pdf(src, output, progress=None) -> tuple[Path, str, str | None]:
    """Devuelve (ruta, método, nota). Intenta PowerPoint, luego LibreOffice y por
    último un renderizado básico con python-pptx + ReportLab."""
    from .conversions import _find_soffice, _soffice_convert

    progress = progress or _noop
    errors = []
    progress(0.1, "Convirtiendo con Microsoft PowerPoint…")
    try:
        _pptx_via_powerpoint(src, output)
        progress(1.0, None)
        return Path(output), "Microsoft PowerPoint", None
    except Exception as exc:
        errors.append(f"PowerPoint: {exc}")
    if _find_soffice():
        progress(0.4, "Convirtiendo con LibreOffice…")
        try:
            _soffice_convert(src, output, "pdf")
            progress(1.0, None)
            return Path(output), "LibreOffice", None
        except Exception as exc:
            errors.append(f"LibreOffice: {exc}")
    progress(0.6, "Renderizando diapositivas…")
    try:
        _pptx_basic(src, output)
    except Exception as exc:
        errors.append(f"básico: {exc}")
        raise RuntimeError("No se pudo convertir la presentación. " + " · ".join(errors)) from exc
    progress(1.0, None)
    return Path(output), "renderizado básico", BASIC_PPTX_NOTE


def _pptx_via_powerpoint(src, output) -> None:
    import pythoncom
    import win32com.client

    pythoncom.CoInitialize()
    try:
        app = win32com.client.DispatchEx("PowerPoint.Application")
        presentation = app.Presentations.Open(str(Path(src).resolve()), True, False, False)
        try:
            presentation.SaveAs(str(Path(output).resolve()), 32)  # 32 = ppSaveAsPDF
        finally:
            presentation.Close()
            # PowerPoint tiene una sola instancia: no cerrar la del usuario si la estaba usando
            if app.Presentations.Count == 0:
                app.Quit()
        if not Path(output).exists():
            raise RuntimeError("PowerPoint no generó el archivo.")
    finally:
        pythoncom.CoUninitialize()


def _pptx_basic(src, output) -> None:
    from pptx import Presentation
    from reportlab.pdfgen import canvas

    prs = Presentation(str(src))
    width, height = prs.slide_width / EMU_PER_PT, prs.slide_height / EMU_PER_PT
    c = canvas.Canvas(str(output), pagesize=(width, height))
    for slide in prs.slides:
        c.setFillColorRGB(1, 1, 1)
        c.rect(0, 0, width, height, stroke=0, fill=1)
        background = _solid_rgb(getattr(slide.background, "fill", None))
        if background:
            c.setFillColorRGB(*background)
            c.rect(0, 0, width, height, stroke=0, fill=1)
        for shape in slide.shapes:
            _draw_shape(c, shape, height)
        c.showPage()
    c.save()


def _solid_rgb(fill):
    from pptx.enum.dml import MSO_FILL

    try:
        if fill is not None and fill.type == MSO_FILL.SOLID:
            rgb = fill.fore_color.rgb
            return tuple(int(str(rgb)[i:i + 2], 16) / 255 for i in (0, 2, 4))
    except Exception:
        pass  # colores del tema u otros tipos de relleno: se omiten
    return None


def _font_rgb(font):
    try:
        rgb = font.color.rgb
        if rgb is not None:
            return tuple(int(str(rgb)[i:i + 2], 16) / 255 for i in (0, 2, 4))
    except Exception:
        pass
    return (0, 0, 0)


def _draw_shape(c, shape, page_h) -> None:
    from pptx.enum.text import PP_ALIGN
    from reportlab.lib.utils import ImageReader, simpleSplit

    if shape.left is None or shape.top is None:
        return
    x, top = shape.left / EMU_PER_PT, shape.top / EMU_PER_PT
    w, h = (shape.width or 0) / EMU_PER_PT, (shape.height or 0) / EMU_PER_PT
    if getattr(shape, "shapes", None) is not None and shape.shape_type == 6:  # grupo
        for child in shape.shapes:
            _draw_shape(c, child, page_h)
        return
    fill = _solid_rgb(getattr(shape, "fill", None))
    if fill and w and h:
        c.setFillColorRGB(*fill)
        c.rect(x, page_h - top - h, w, h, stroke=0, fill=1)
    if hasattr(shape, "image"):
        try:
            c.drawImage(ImageReader(BytesIO(shape.image.blob)), x, page_h - top - h, w, h,
                        mask="auto")
        except Exception:
            pass  # EMF/WMF y otros formatos que ReportLab no lee
    if getattr(shape, "has_table", False) and shape.has_table:
        _draw_table(c, shape, x, top, page_h)
        return
    if not (getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip()):
        return
    cursor = top + 7.2
    inner_w = max(10, w - 14.4)
    for paragraph in shape.text_frame.paragraphs:
        runs = paragraph.runs
        text = "".join(r.text for r in runs)
        font = runs[0].font if runs else paragraph.font
        size = font.size.pt if font.size else 18
        name = "Helvetica-Bold" if font.bold else "Helvetica"
        c.setFillColorRGB(*_font_rgb(font))
        c.setFont(name, size)
        for line in simpleSplit(text, name, size, inner_w) or [""]:
            cursor += size * 1.15
            baseline = page_h - cursor
            if paragraph.alignment == PP_ALIGN.CENTER:
                c.drawCentredString(x + w / 2, baseline, line)
            elif paragraph.alignment == PP_ALIGN.RIGHT:
                c.drawRightString(x + w - 7.2, baseline, line)
            else:
                c.drawString(x + 7.2, baseline, line)


def _draw_table(c, shape, x, top, page_h) -> None:
    from reportlab.lib.utils import simpleSplit

    table = shape.table
    cell_y = top
    for row in table.rows:
        row_h = row.height / EMU_PER_PT
        cell_x = x
        for col, cell in zip(table.columns, row.cells):
            col_w = col.width / EMU_PER_PT
            c.setStrokeColorRGB(0.6, 0.6, 0.6)
            c.rect(cell_x, page_h - cell_y - row_h, col_w, row_h, stroke=1, fill=0)
            c.setFillColorRGB(0, 0, 0)
            c.setFont("Helvetica", 10)
            for n, line in enumerate(simpleSplit(cell.text, "Helvetica", 10, col_w - 6)[:4]):
                c.drawString(cell_x + 3, page_h - cell_y - 12 - n * 12, line)
            cell_x += col_w
        cell_y += row_h


# --------------------------------------------------------------------------- HTML -> PDF

def _find_edge() -> str | None:
    for candidate in (
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
    ):
        if os.path.exists(candidate):
            return candidate
    return shutil.which("msedge")


def html_to_pdf(source: str, output, progress=None) -> tuple[Path, str]:
    """`source` = ruta a un .html o una URL http(s). Devuelve (ruta, motor usado).
    Usa WeasyPrint (incluido) y, si falla, Microsoft Edge en modo headless."""
    progress = progress or _noop
    source = source.strip()
    is_url = source.lower().startswith(("http://", "https://"))
    if not is_url and not Path(source).is_file():
        raise ValueError("No se encontró el archivo HTML.")
    errors = []

    progress(0.1, "Convirtiendo con WeasyPrint…")
    bundled = assets_dir() / "weasyprint" / "weasyprint.exe"
    try:
        if bundled.exists():
            proc = _run([str(bundled), source, str(output)])
            if proc.returncode != 0 or not Path(output).exists():
                raise RuntimeError(_last_error_line(proc))
        else:
            import weasyprint  # requiere GTK/Pango instalados en el sistema
            weasyprint.HTML(source if is_url else Path(source).resolve().as_uri()).write_pdf(str(output))
        progress(1.0, None)
        return Path(output), "WeasyPrint"
    except Exception as exc:
        errors.append(f"WeasyPrint: {exc}")

    edge = _find_edge()
    if edge:
        progress(0.5, "Reintentando con Microsoft Edge…")
        target = source if is_url else Path(source).resolve().as_uri()
        with tempfile.TemporaryDirectory() as profile:
            proc = _run([edge, "--headless", "--disable-gpu", "--no-first-run",
                         "--no-pdf-header-footer", f"--user-data-dir={profile}",
                         f"--print-to-pdf={Path(output).resolve()}", target], timeout=180)
        if Path(output).exists() and Path(output).stat().st_size > 0:
            progress(1.0, None)
            return Path(output), "Microsoft Edge"
        errors.append(f"Edge: {_last_error_line(proc)}")
    hint = " Revisa la conexión a internet." if is_url else ""
    raise RuntimeError("No se pudo convertir el HTML a PDF." + hint + " (" + " · ".join(errors) + ")")


# --------------------------------------------------------------------------- PDF -> PowerPoint

_PPTX_MAX_PT = 56 * 72  # límite de PowerPoint: 56 pulgadas por lado


def pdf_to_pptx(src, output, dpi: int = 150, progress=None) -> tuple[Path, int]:
    """Cada página pasa a ser una diapositiva con la página como imagen de fondo."""
    from pptx import Presentation
    from pptx.util import Emu

    from .page_tools import page_sizes_pt
    from .pdf_tools import render_page_image

    progress = progress or _noop
    _open(src)
    sizes = page_sizes_pt(src)
    slide_w, slide_h = sizes[0]
    factor = min(1.0, _PPTX_MAX_PT / max(slide_w, slide_h))
    slide_w, slide_h = slide_w * factor, slide_h * factor
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(int(slide_w * EMU_PER_PT)), Emu(int(slide_h * EMU_PER_PT))
    blank = prs.slide_layouts[6]
    for i, (page_w, page_h) in enumerate(sizes):
        progress(i / len(sizes), f"Página {i + 1} de {len(sizes)}…")
        image = render_page_image(src, i, dpi=dpi).convert("RGB")
        buf = BytesIO()
        image.save(buf, "JPEG", quality=88)
        buf.seek(0)
        scale = min(slide_w / page_w, slide_h / page_h)
        pic_w, pic_h = page_w * scale, page_h * scale
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(buf, Emu(int((slide_w - pic_w) / 2 * EMU_PER_PT)),
                                 Emu(int((slide_h - pic_h) / 2 * EMU_PER_PT)),
                                 Emu(int(pic_w * EMU_PER_PT)), Emu(int(pic_h * EMU_PER_PT)))
    prs.save(str(output))
    progress(1.0, None)
    return Path(output), len(sizes)


# --------------------------------------------------------------------------- PDF -> PDF/A

LEVEL_FULL = "PDF/A-1b completo (Ghostscript; el texto se conserva)"
LEVEL_BASIC = ("PDF/A-1b básico (las páginas se guardaron como imagen: el texto ya no "
               "se puede seleccionar; instala Ghostscript para conservarlo)")


def find_ghostscript() -> str | None:
    for name in ("gswin64c", "gswin32c", "gs"):
        found = shutil.which(name)
        if found:
            return found
    for root in (r"C:\Program Files\gs", r"C:\Program Files (x86)\gs"):
        matches = sorted(glob.glob(os.path.join(root, "gs*", "bin", "gswin*c.exe")))
        if matches:
            return matches[-1]
    return None


def pdf_to_pdfa(src, output, progress=None) -> tuple[Path, str]:
    """Devuelve (ruta, nivel de cumplimiento alcanzado)."""
    progress = progress or _noop
    _open(src)
    gs = find_ghostscript()
    if gs:
        progress(0.2, "Convirtiendo con Ghostscript…")
        proc = _run([gs, "-dPDFA=1", "-dBATCH", "-dNOPAUSE", "-dNOSAFER",
                     "-sColorConversionStrategy=UseDeviceIndependentColor",
                     "-sDEVICE=pdfwrite", "-dPDFACompatibilityPolicy=1",
                     f"-sOutputFile={output}", str(src)], timeout=600)
        if proc.returncode == 0 and Path(output).exists():
            progress(1.0, None)
            return Path(output), LEVEL_FULL
    _pdfa_raster(src, output, progress)
    progress(1.0, None)
    return Path(output), LEVEL_BASIC


def _pdfa_raster(src, output, progress, dpi: int = 200) -> None:
    import img2pdf
    import pikepdf
    from PIL import ImageCms

    from .pdf_tools import get_page_count, render_page_image

    total = get_page_count(src)
    pages = []
    for i in range(total):
        progress(0.9 * i / total, f"Página {i + 1} de {total}…")
        image = render_page_image(src, i, dpi=dpi).convert("RGB")  # PDF/A-1 no admite transparencias
        buf = BytesIO()
        image.save(buf, "JPEG", quality=90, dpi=(dpi, dpi))
        pages.append(buf.getvalue())
    progress(0.95, "Agregando metadatos PDF/A…")
    with tempfile.TemporaryDirectory() as tmp:
        icc = Path(tmp) / "sRGB.icc"
        icc.write_bytes(ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB")).tobytes())
        data = img2pdf.convert(pages, pdfa=str(icc), title=Path(src).stem)
    # PDF/A-1 se basa en PDF 1.4 y no admite flujos de objetos
    with pikepdf.open(BytesIO(data)) as pdf:
        pdf.save(str(output), force_version="1.4",
                 object_stream_mode=pikepdf.ObjectStreamMode.disable)
