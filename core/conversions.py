"""Document format conversions (PDF <-> Word/Excel). No UI imports.

External programs used as engines/fallbacks when available:
  - Microsoft Word (via docx2pdf/COM) for Word -> PDF
  - LibreOffice headless (soffice) as fallback for Word -> PDF and Excel -> PDF
"""
from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path


def _noop(fraction: float, message: str | None = None) -> None:
    pass


# --------------------------------------------------------------------------- helpers

def _find_soffice() -> str | None:
    exe = shutil.which("soffice")
    if exe:
        return exe
    for candidate in (
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
        "/usr/bin/soffice",
        "/usr/local/bin/soffice",
        "/Applications/LibreOffice.app/Contents/MacOS/soffice",
    ):
        if os.path.exists(candidate):
            return candidate
    return None


def _soffice_convert(src: str | Path, output: str | Path, target: str = "pdf") -> None:
    exe = _find_soffice()
    if not exe:
        raise RuntimeError("LibreOffice is not installed.")
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(
            [exe, "--headless", "--convert-to", target, "--outdir", tmp, str(src)],
            check=True,
            timeout=300,
            capture_output=True,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
        )
        produced = Path(tmp) / (Path(src).stem + f".{target}")
        if not produced.exists():
            raise RuntimeError("LibreOffice did not produce an output file.")
        shutil.move(str(produced), str(output))


# --------------------------------------------------------------------------- PDF -> Word

def pdf_to_word(src: str | Path, output: str | Path, progress=None) -> Path:
    from pdf2docx import Converter

    progress = progress or _noop
    progress(0.05, "Analyzing PDF layout…")
    converter = Converter(str(src))
    try:
        converter.convert(str(output))
    except Exception as exc:
        raise RuntimeError(
            f"Could not convert this PDF to Word: {exc}. Scanned/image-only PDFs "
            "cannot be converted without OCR."
        ) from exc
    finally:
        converter.close()
    progress(1.0, None)
    return Path(output)


# --------------------------------------------------------------------------- Word -> PDF

def word_to_pdf(src: str | Path, output: str | Path, progress=None) -> Path:
    progress = progress or _noop
    progress(0.1, "Converting with Microsoft Word…")
    word_error: Exception | None = None
    try:
        _word_to_pdf_msword(src, output)
        progress(1.0, None)
        return Path(output)
    except Exception as exc:
        word_error = exc

    progress(0.5, "Word unavailable — trying LibreOffice…")
    try:
        _soffice_convert(src, output, "pdf")
    except Exception as lo_error:
        raise RuntimeError(
            "Word → PDF needs Microsoft Word or LibreOffice installed; neither "
            f"worked. (Word: {word_error}; LibreOffice: {lo_error})"
        ) from lo_error
    progress(1.0, None)
    return Path(output)


def _word_to_pdf_msword(src: str | Path, output: str | Path) -> None:
    # docx2pdf drives Word over COM; COM must be initialized in this worker thread.
    try:
        import pythoncom
    except ImportError:
        pythoncom = None
    if pythoncom:
        pythoncom.CoInitialize()
    try:
        from docx2pdf import convert
        convert(str(src), str(output))
        if not Path(output).exists():
            raise RuntimeError("Word did not produce an output file.")
    finally:
        if pythoncom:
            pythoncom.CoUninitialize()


# --------------------------------------------------------------------------- PDF -> Excel

def pdf_to_excel(src: str | Path, output: str | Path, progress=None) -> tuple[Path, int]:
    """Extract tables into an .xlsx workbook (one sheet per table).

    Returns (output_path, table_count). Tries camelot (stream, then lattice),
    then tabula-py if installed.
    """
    progress = progress or _noop
    tables = None
    errors: list[str] = []

    try:
        import camelot
        progress(0.1, "Detecting tables (stream mode)…")
        found = camelot.read_pdf(str(src), pages="all", flavor="stream")
        if len(found) == 0:
            progress(0.4, "Retrying in lattice mode…")
            try:
                found = camelot.read_pdf(str(src), pages="all", flavor="lattice")
            except Exception as exc:
                errors.append(f"camelot lattice: {exc}")
        if len(found):
            tables = [(t.df, getattr(t, "page", i + 1)) for i, t in enumerate(found)]
    except ImportError:
        errors.append("camelot-py is not installed")
    except Exception as exc:
        errors.append(f"camelot: {exc}")

    if tables is None:
        try:
            import tabula
            progress(0.5, "Trying tabula-py…")
            dfs = tabula.read_pdf(str(src), pages="all", multiple_tables=True)
            if dfs:
                tables = [(df, i + 1) for i, df in enumerate(dfs)]
        except ImportError:
            errors.append("tabula-py is not installed")
        except Exception as exc:
            errors.append(f"tabula: {exc}")

    if not tables:
        detail = "; ".join(errors) if errors else "no tables detected"
        raise RuntimeError(
            f"No tables could be extracted from this PDF ({detail}). "
            "This works best on PDFs with real (text-based) tables."
        )

    progress(0.8, f"Writing {len(tables)} table(s) to Excel…")
    import pandas as pd
    with pd.ExcelWriter(str(output), engine="openpyxl") as writer:
        for i, (df, page) in enumerate(tables):
            sheet = f"Table {i + 1} (p{page})"[:31]
            df.to_excel(writer, sheet_name=sheet, index=False, header=False)
    progress(1.0, None)
    return Path(output), len(tables)


# --------------------------------------------------------------------------- Excel -> PDF

def excel_to_pdf(src: str | Path, output: str | Path, progress=None) -> Path:
    progress = progress or _noop
    try:
        progress(0.1, "Rendering with ReportLab…")
        _excel_to_pdf_reportlab(src, output)
        progress(1.0, None)
        return Path(output)
    except Exception as rl_error:
        progress(0.5, "Trying LibreOffice fallback…")
        try:
            _soffice_convert(src, output, "pdf")
        except Exception:
            raise RuntimeError(
                f"Could not render this workbook to PDF: {rl_error}. Installing "
                "LibreOffice enables a higher-fidelity fallback."
            ) from rl_error
        progress(1.0, None)
        return Path(output)


def _excel_to_pdf_reportlab(src: str | Path, output: str | Path) -> None:
    from openpyxl import load_workbook
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    workbook = load_workbook(str(src), data_only=True, read_only=True)
    styles = getSampleStyleSheet()
    elements = []
    for sheet in workbook.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if v is None else str(v) for v in row]
            rows.append([c if len(c) <= 80 else c[:77] + "…" for c in cells])
        # drop fully empty trailing rows/columns
        while rows and not any(rows[-1]):
            rows.pop()
        if rows:
            width = max(len(r) for r in rows)
            while width and not any(r[width - 1] if len(r) >= width else "" for r in rows):
                width -= 1
            rows = [(r + [""] * width)[:width] for r in rows]
        if not rows:
            continue
        elements.append(Paragraph(sheet.title, styles["Heading2"]))
        elements.append(Spacer(0, 4 * mm))
        table = Table(rows, repeatRows=1)
        table.setStyle(TableStyle([
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e8eef7")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]))
        elements.append(table)
        elements.append(PageBreak())
    workbook.close()
    if not elements:
        raise RuntimeError("The workbook has no data to render.")
    doc = SimpleDocTemplate(
        str(output), pagesize=landscape(A4),
        leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
        title=Path(src).stem,
    )
    doc.build(elements[:-1] if isinstance(elements[-1], PageBreak) else elements)
